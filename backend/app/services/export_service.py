import io
import logging
from typing import List, Any
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# Configurar el registrador de eventos (logger)
logger = logging.getLogger(__name__)

def generar_excel(elementos: List[Any]) -> bytes:
    """
    Genera un archivo Excel (.xlsx) altamente formateado y profesional con el cronometraje de los elementos.
    Contiene dos hojas:
    - "Cronometraje": Muestra los elementos agrupados con su categoría, duración promedio y número de muestras.
    - "Detalle Muestras": Muestra el desglose detallado de cada muestra observada (inicio en video y duración).

    Args:
        elementos (List[Any]): Lista de diccionarios o modelos de Pydantic con los elementos de trabajo.

    Returns:
        bytes: El contenido del archivo Excel en bytes listo para ser transmitido o descargado.
    """
    try:
        logger.info(f"Iniciando la generación del archivo Excel multioja para {len(elementos)} elementos.")
        
        # Crear un nuevo libro
        wb = Workbook()
        
        # -------------------------------------------------------------
        # HOJA 1: Cronometraje
        # -------------------------------------------------------------
        ws = wb.active
        ws.title = "Cronometraje"
        ws.views.sheetView[0].showGridLines = True
        
        # Cabeceras
        headers = ["N°", "Elemento", "Categoría", "Duración", "Muestras"]
        ws.append(headers)
        
        # Estilos generales elegantes
        font_header = Font(name="Calibri", size=11, bold=True, color="08060D")
        fill_header = PatternFill(start_color="E5E4E7", end_color="E5E4E7", fill_type="solid") # Gris suave
        
        align_center = Alignment(horizontal="center", vertical="center")
        align_left = Alignment(horizontal="left", vertical="center")
        align_right = Alignment(horizontal="right", vertical="center")
        
        # Bordes finos elegantes para todas las celdas
        border_side = Side(style='thin', color='B0B0B0')
        border_all = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)
        
        # Insertar datos de los elementos
        for elem in elementos:
            # Soportar tanto objetos Pydantic como diccionarios estándar
            if hasattr(elem, "model_dump"):
                e = elem.model_dump()
            elif hasattr(elem, "dict"):
                e = elem.dict()
            elif hasattr(elem, "__dict__"):
                e = elem.__dict__
            else:
                e = elem
                
            muestras_list = e.get("muestras", [])
            duracion_prom = float(e.get("duracion_promedio", 0.0))
            num_muestras = len(muestras_list)
            
            row_data = [
                int(e.get("numero", 0)),
                str(e.get("elemento", "")),
                str(e.get("categoria", "Tmanual")),
                duracion_prom,
                num_muestras
            ]
            ws.append(row_data)
            
        # Insertar fila de Totales
        tot_row_idx = len(elementos) + 2
        ws.cell(row=tot_row_idx, column=1, value="Total").font = Font(name="Calibri", size=11, bold=True)
        ws.cell(row=tot_row_idx, column=1).alignment = align_left
        
        # Calcular suma total de duraciones promedio y total de muestras
        total_duracion = 0.0
        total_muestras_count = 0
        for elem in elementos:
            if hasattr(elem, "model_dump"):
                e = elem.model_dump()
            elif hasattr(elem, "dict"):
                e = elem.dict()
            elif hasattr(elem, "__dict__"):
                e = elem.__dict__
            else:
                e = elem
            total_duracion += float(e.get("duracion_promedio", 0.0))
            total_muestras_count += len(e.get("muestras", []))
            
        ws.cell(row=tot_row_idx, column=4, value=total_duracion).font = Font(name="Calibri", size=11, bold=True)
        ws.cell(row=tot_row_idx, column=4).alignment = align_right
        
        ws.cell(row=tot_row_idx, column=5, value=total_muestras_count).font = Font(name="Calibri", size=11, bold=True)
        ws.cell(row=tot_row_idx, column=5).alignment = align_right
        
        # Aplicar formatos de celdas, estilos y bordes para la Hoja 1
        for r_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=tot_row_idx, min_col=1, max_col=5), start=1):
            for c_idx, cell in enumerate(row, start=1):
                cell.border = border_all
                
                if r_idx == 1:
                    cell.font = font_header
                    cell.fill = fill_header
                    cell.alignment = align_center
                elif r_idx == tot_row_idx:
                    cell.font = Font(name="Calibri", size=11, bold=True)
                    if c_idx == 1:
                        cell.alignment = align_left
                    elif c_idx == 4:
                        cell.alignment = align_right
                        cell.number_format = '0.0"s"'
                    elif c_idx == 5:
                        cell.alignment = align_right
                        cell.number_format = '#,##0'
                else:
                    if c_idx == 1:
                        cell.alignment = align_center
                    elif c_idx in (2, 3):
                        cell.alignment = align_left
                    elif c_idx == 4:
                        cell.alignment = align_right
                        cell.number_format = '0.0"s"'
                    elif c_idx == 5:
                        cell.alignment = align_right
                        cell.number_format = '#,##0'
                        
        # Auto-ajustar el ancho de las columnas Hoja 1
        for col in ws.columns:
            max_len = 0
            col_letter = col[0].column_letter
            for cell in col:
                val_str = str(cell.value or '')
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 11)
            
        # -------------------------------------------------------------
        # HOJA 2: Detalle Muestras
        # -------------------------------------------------------------
        ws2 = wb.create_sheet(title="Detalle Muestras")
        ws2.views.sheetView[0].showGridLines = True
        
        # Cabeceras Hoja 2
        headers2 = ["Elemento", "N° Muestra", "Inicio Video (seg)", "Duración (seg)"]
        ws2.append(headers2)
        
        # Rellenar datos Hoja 2
        total_muestras = 0
        for elem in elementos:
            if hasattr(elem, "model_dump"):
                e = elem.model_dump()
            elif hasattr(elem, "dict"):
                e = elem.dict()
            elif hasattr(elem, "__dict__"):
                e = elem.__dict__
            else:
                e = elem
                
            elemento_nombre = str(e.get("elemento", ""))
            muestras_list = e.get("muestras", [])
            for m in muestras_list:
                row_data2 = [
                    elemento_nombre,
                    int(m.get("numero", 0)),
                    float(m.get("inicio_video", 0.0)),
                    float(m.get("duracion", 0.0))
                ]
                ws2.append(row_data2)
                total_muestras += 1
                
        # Aplicar formatos de celdas, estilos y bordes para la Hoja 2
        for r_idx, row in enumerate(ws2.iter_rows(min_row=1, max_row=total_muestras + 1, min_col=1, max_col=4), start=1):
            for c_idx, cell in enumerate(row, start=1):
                cell.border = border_all
                
                if r_idx == 1:
                    cell.font = font_header
                    cell.fill = fill_header
                    cell.alignment = align_center
                else:
                    if c_idx == 1:
                        cell.alignment = align_left
                    elif c_idx == 2:
                        cell.alignment = align_center
                    else:
                        cell.alignment = align_right
                        cell.number_format = '0.0"s"'
                        
        # Auto-ajustar el ancho de las columnas Hoja 2
        for col in ws2.columns:
            max_len = 0
            col_letter = col[0].column_letter
            for cell in col:
                val_str = str(cell.value or '')
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws2.column_dimensions[col_letter].width = max(max_len + 4, 11)
            
        # -------------------------------------------------------------
        # GUARDAR Y RETORNAR BYTES
        # -------------------------------------------------------------
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        
        logger.info("Archivo Excel multioja generado con éxito.")
        return buffer.getvalue()

    except Exception as e:
        logger.error(f"Error crítico al generar el archivo Excel: {str(e)}")
        raise e
